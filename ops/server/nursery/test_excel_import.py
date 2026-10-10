# -*- coding: utf-8 -*-
"""اختبارات قارئ Excel الشهري — تشتغل بدون Odoo.

تشغيل:  python3 ops/server/nursery/test_excel_import.py

الشيت المصطنع هنا بنفس تخطيط ملف الحسابات الحقيقي (طلاب بأرقام تسلسلية
كمعادلات، جدول رواتب وجدول مصروفات جنب بعض، صف ملخص فوق) ومن غير أي
بيانات حقيقية.
"""
import base64
import importlib.util
import io
import json
import os
import sys
import types
import unittest

from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))


def _load_module():
    """Import excel_import.py with the Odoo imports stubbed out."""
    odoo = types.ModuleType('odoo')
    odoo_http = types.ModuleType('odoo.http')
    odoo_http.request = None
    odoo_http.Controller = object
    odoo_http.route = lambda *args, **kwargs: (lambda func: func)
    odoo.http = odoo_http
    odoo.fields = types.SimpleNamespace()
    sys.modules.setdefault('odoo', odoo)
    sys.modules.setdefault('odoo.http', odoo_http)
    try:
        import pytz  # noqa: F401 — used only inside the controller routes
    except ImportError:
        sys.modules.setdefault('pytz', types.ModuleType('pytz'))
    package = types.ModuleType('nursery_test_pkg')
    package.__path__ = []
    roles = types.ModuleType('nursery_test_pkg.roles')
    roles.WEBSITE = ''
    roles._manager_by_token = lambda *args, **kwargs: None
    sys.modules['nursery_test_pkg'] = package
    sys.modules['nursery_test_pkg.roles'] = roles
    spec = importlib.util.spec_from_file_location(
        'nursery_test_pkg.excel_import', os.path.join(HERE, 'excel_import.py'))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


excel_import = _load_module()

REAL_WORKBOOK = os.path.join(HERE, 'randa.xlsx')

# الـ15 خلية في كتلة الملخص بترتيب كروت اللوحة (عمود عمود): القيم من الشيت الحقيقي
# (CONTRACT.md) لسبتمبر وأكتوبر 2026، ونوفمبر شهر فاضي.
CONTRACT_KEYS = ('opening_balance', 'collected', 'randa_cash', 'bank_transfer', 'expenses',
                 'on_hand_randa', 'expected_salaries', 'cash_after_salaries', 'salaries_paid',
                 'net', 'student_count', 'paid_students', 'unpaid_students', 'books',
                 'remaining_total')
CONTRACT_VALUES = {
    '2026-09': (0, 276215, 234215, 42000, 230125, 4090, 39750, 4090, 36450, 46090,
                87, 50, 0, 54000, 14265),
    '2026-10': (4090, 77750, 77750, 0, 1450, 76300, 39750, 36550, 0, 76300,
                86, 80, 6, 5700, 11100),
    '2026-11': (0, 0, 0, 0, 0, 0, 39750, -39750, 0, 0, 86, 0, 0, 0, 102100),
}

STUDENT_HEADERS = ['Class', 'Serial', 'Student', 'Joining Date', 'Agreed Monthly Fees',
                   'Last Payment Amount', 'Payment Date', 'Paid Until Date', 'Books',
                   'Fees paid', 'Remaining', 'Payment status', 'Payment Method',
                   'ملاحظات', 'مرتبط بولي الأمر', 'اسم ولي الأمر', 'هاتف ولي الأمر', 'ID']


def _workbook():
    """Two classes, one mid-month leaver, salaries + expenses side by side."""
    book = Workbook()
    sheet = book.active
    sheet.title = 'September 2026'
    sheet['A1'] = 'Accounting Summary'
    for column, label in zip('ACEGIKMOQS', [
            'Total Received', 'Randa Cash', 'Bank Transfer', 'Expenses', 'On hand Randa',
            'Salaries', 'Students\nعدد الطلاب', 'Books Fees\nرسوم الكتب',
            'Remaining\nالباقي', 'Net']):
        sheet['%s2' % column] = label
    sheet.append([])
    sheet.append([])
    sheet.append(STUDENT_HEADERS)
    rows = [
        # class, name, joined, fees, paid, paid_until, books, fees_paid, method, id, serial?
        ('KG1', 'Alpha', '2026-09-01', 1000, 8000, '2027-05-30', 1000, 7000, 'Randa Cash', 11, True),
        ('KG1', 'Beta', '2026-09-01', 1200, 2300, '2026-10-01', 1000, 1300, 'Randa Cash', 12, True),
        ('KG1', 'Gone', '2026-09-01', 1300, 1300, '2026-10-01', 500, 1800, 'Randa Cash', 13, False),
        ('KG1', 'Gamma', '2026-09-27', 1300, None, None, None, None, None, 14, True),
        ('KG2', 'Delta', '2026-09-01', 1000, 10500, '2027-05-30', 500, 10000, 'Bank Transfer', 21, True),
        ('KG2', 'Epsilon', '2026-09-01', 1100, 600, '2026-10-01', None, 600, 'Randa Cash', 22, True),
    ]
    class_start = {}
    for cls, name, joined, fees, paid, until, books, fees_paid, method, db_id, numbered in rows:
        row_number = sheet.max_row + 1
        start = class_start.setdefault(cls, row_number)
        serial = ('=IF(C{r}="","",COUNT(B${s}:B{p})+1)'.format(r=row_number, s=start, p=row_number - 1)
                  if row_number > start else '=IF(C{r}="","",1)'.format(r=row_number))
        sheet.append([
            cls if row_number == start else None,
            serial if numbered else None,
            name, joined, fees, paid, joined if paid else None, until, books, fees_paid,
            '=IF(OR(C{r}="",E{r}=""),"",MAX(0,E{r}-J{r}))'.format(r=row_number),
            'paid' if paid else None, method, None, 'لا', None, None, db_id,
        ])
    # one blank template row with formulas only, then the class counter rows
    row_number = sheet.max_row + 1
    sheet.append([None, '=IF(C{r}="","",COUNT(B$10:B{p})+1)'.format(r=row_number, p=row_number - 1),
                  None, None, None, None, None, None, None, None,
                  '=IF(OR(C{r}="",E{r}=""),"",MAX(0,E{r}-J{r}))'.format(r=row_number)])
    last_student = sheet.max_row
    sheet.append([None, None, 'Students KG2', '=COUNT(B10:B%d)' % last_student])
    kg2_count_row = sheet.max_row
    sheet.append([None, None, 'Students KG1', '=COUNT(B6:B9)'])
    kg1_count_row = sheet.max_row
    sheet.append([])

    sheet.append(['Salaries\nالرواتب', None, None, None, None, None, 'Expenses Paid by Randa'])
    sheet.append(['Teacher Name\nاسم المعلم', 'Date\nالتاريخ', 'Forecast Salary', 'Actual Salary',
                  'Notes', None, 'Serial', 'Expense Name', 'Expense Date', 'Value'])
    first_entry = sheet.max_row + 1
    entries = [
        (('Dalia', None, 4000, 4000, None), ('نسرين', '2026-08-28', 10360)),
        (('Randa', None, 3500, 2500, 'كتب'), ('قرطاسيه', '2026-08-29', -475)),
        (('Huda', None, 2200, 2200, None), ('مرتبات', None, -8700)),
        ((None, None, None, None, None), ('ايجار', '2026-09-10', -20000)),
    ]
    for salary, expense in entries:
        row_number = sheet.max_row + 1
        serial = ('=IF(H{r}="","",1)'.format(r=row_number) if row_number == first_entry
                  else '=IF(H{r}="","",COUNT(G${s}:G{p})+1)'.format(r=row_number, s=first_entry,
                                                                   p=row_number - 1))
        sheet.append(list(salary) + [None, serial] + list(expense))
    last_entry = sheet.max_row
    sheet.append(['Total Salaries\nإجمالي الرواتب', None,
                  '=SUM(C%d:C%d)' % (first_entry, last_entry),
                  '=SUM(D%d:D%d)' % (first_entry, last_entry), None, None, None,
                  'Total Expenses', None, '=SUM(J%d:J%d)' % (first_entry, last_entry)])
    total_row = sheet.max_row

    sheet['A3'] = '=SUM(F6:F%d)' % last_student
    sheet['C3'] = '=SUMIF(M6:M%d,"Randa Cash",F6:F%d)' % (last_student, last_student)
    sheet['G3'] = '=-J%d' % total_row
    sheet['I3'] = '=C3-G3'
    sheet['K3'] = '=D%d' % total_row
    sheet['M3'] = '=D%d+D%d' % (kg1_count_row, kg2_count_row)
    sheet['O3'] = '=SUM(I6:I%d)' % last_student
    sheet['Q3'] = '=SUM(K6:K%d)' % last_student
    sheet['S3'] = '=A3-G3'
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


class ExcelImportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        raw = _workbook()
        cls.parsed = excel_import._parse_records(
            [{'name': 'sep.xlsx', 'content': base64.b64encode(raw).decode('ascii')}],
            '2026-09')
        cls.metrics = excel_import._excel_month_metrics(cls.parsed, '2026-09')

    def test_workbook_months_prefers_newest_finished_tab(self):
        files = [{'name': 'sep.xlsx', 'content': base64.b64encode(_workbook()).decode('ascii')}]
        self.assertEqual(excel_import._workbook_months(files, '2026-10'), ['2026-09'])
        self.assertEqual(excel_import._workbook_months(files, '2026-09'), ['2026-09'])
        # a tab for a month that has not started yet is never picked
        self.assertEqual(excel_import._workbook_months(files, '2026-08'), [])

    def test_import_unlock_window(self):
        import datetime
        import json, tempfile
        original = excel_import.UNLOCK_FILE
        # The shipped import_unlock.json holds no window (September 2026 is
        # final since 2026-10-03), so the window semantics are tested on a
        # temporary file; the shipped file must keep every month locked.
        excel_import.UNLOCK_FILE = os.path.join(HERE, 'import_unlock.json')
        try:
            self.assertFalse(excel_import._import_unlocked('2026-09', datetime.date(2026, 10, 2)))
        finally:
            excel_import.UNLOCK_FILE = original
        with tempfile.NamedTemporaryFile('w', suffix='.json', delete=False) as handle:
            json.dump({'2026-09': '2026-10-10'}, handle)
        excel_import.UNLOCK_FILE = handle.name
        try:
            self.assertTrue(excel_import._import_unlocked('2026-09', datetime.date(2026, 10, 10)))
            self.assertFalse(excel_import._import_unlocked('2026-09', datetime.date(2026, 10, 11)))
            self.assertFalse(excel_import._import_unlocked('2026-10', datetime.date(2026, 10, 2)))
        finally:
            excel_import.UNLOCK_FILE = original
            os.unlink(handle.name)
        excel_import.UNLOCK_FILE = os.path.join(HERE, 'does-not-exist.json')
        try:
            self.assertFalse(excel_import._import_unlocked('2026-09', datetime.date(2026, 10, 2)))
        finally:
            excel_import.UNLOCK_FILE = original

    def test_formula_serials_are_resolved(self):
        """openpyxl saves formulas without cached values; serials still count."""
        by_name = {row['name']: row for row in self.parsed['students']}
        self.assertEqual([by_name[n]['serial'] for n in ('Alpha', 'Beta', 'Gamma')], [1, 2, 3])
        self.assertEqual([by_name[n]['serial'] for n in ('Delta', 'Epsilon')], [1, 2])
        self.assertIsNone(by_name['Gone']['serial'])
        self.assertFalse(by_name['Gone']['active'])
        self.assertEqual(self.metrics['student_count'], 5)

    def test_remaining_formula_treats_blank_paid_as_zero(self):
        by_name = {row['name']: row for row in self.parsed['students']}
        self.assertEqual(by_name['Gamma']['remaining'], 1300)
        self.assertEqual(by_name['Epsilon']['remaining'], 500)
        self.assertEqual(by_name['Alpha']['remaining'], 0)
        self.assertEqual(self.metrics['remaining_total'], 1800)

    def test_staff_rows_are_not_students(self):
        names = {row['name'] for row in self.parsed['students']}
        self.assertNotIn('Dalia', names)
        self.assertEqual(len(self.parsed['students']), 6)

    def test_salary_and_expense_tables_are_split(self):
        salaries = [row for row in self.parsed['entries'] if row['kind'] == 'salary']
        expenses = [row for row in self.parsed['entries'] if row['kind'] == 'expense']
        self.assertEqual([(r['name'], r['amount'], r['expected']) for r in salaries],
                         [('Dalia', 4000, 4000), ('Randa', 2500, 3500), ('Huda', 2200, 2200)])
        self.assertEqual([(r['name'], r['amount']) for r in expenses],
                         [('نسرين', 10360), ('قرطاسيه', -475), ('مرتبات', -8700), ('ايجار', -20000)])
        self.assertEqual(expenses[0]['date'], '2026-08-28')

    def test_metrics_match_the_workbook_arithmetic(self):
        m = self.metrics
        self.assertEqual(m['collected'], 22700)
        self.assertEqual(m['randa_cash'], 12200)
        self.assertEqual(m['bank_transfer'], 10500)
        self.assertEqual(m['salaries'], 8700)
        self.assertEqual(m['expected_salaries'], 9700)
        self.assertEqual(m['expenses'], 18815)           # -(10360 - 475 - 8700 - 20000)
        self.assertEqual(m['opening_balance'], 10360)
        self.assertEqual(m['salary_payout'], 8700)
        self.assertEqual(m['other_expenses'], 20475)     # 475 + 20000
        self.assertEqual(m['on_hand_randa'], 12200 + 10360 - 29175)
        self.assertEqual(m['net'], 22700 - 18815)
        # dashboard identity: income - salaries - other == workbook net
        self.assertEqual((m['collected'] + m['opening_balance']) - m['salaries'] - m['other_expenses'],
                         m['net'])
        self.assertEqual(m['books'], 3000)


def _summary_block_rows():
    """The real workbook's 12-row summary block (two-line labels) followed by
    the traps that used to poison it: a student header row, data rows with
    'Randa Cash' in the method column, the side-table titles/headers and the
    totals row whose value cell below is empty."""
    rows = [
        ['Accounting Summary — October 2026'],
        ['Opening Balance\nرصيد افتتاحي', None, None, 'On hand Randa\nرصيد راندا',
         None, None, 'Students\nعدد الطلاب'],
        [4090, None, None, 76300, None, None, 86],
        ['Total Received\nإجمالي المحصّل', None, None, 'Expected Salaries\nالرواتب المتوقعة',
         None, None, 'Paid Students\nدافعين'],
        [77750, None, None, 39750, None, None, 80],
        ['Randa Cash\nكاش راندا', None, None, 'Net On hand Randa\nبعد الرواتب',
         None, None, 'Unpaid Students\nغير دافعين'],
        [77750, None, None, 36550, None, None, 6],
        ['Bank Transfer\nتحويل بنكي', None, None, 'Salaries Paid\nرواتب مصروفة',
         None, None, 'Books Fees\nرسوم الكتب'],
        [0, None, None, 0, None, None, 5700],
        ['Expenses\nالمصروفات', None, None, 'Net\nالصافي', None, None, 'Remaining\nالباقي'],
        [1450, None, None, 76300, None, None, 11100],
        [],
        STUDENT_HEADERS,
        ['KG1', 1, 'A', None, 1000, 2500, None, None, None, 2500, 0, 'paid', 'Randa Cash'],
        ['KG1', 2, 'B', None, 1000, None, None, None, None, None, 1000, 'unpaid', None],
        [],
        ['Salaries\nالرواتب', None, None, None, None, None, 'Expenses Paid by Randa\nمصروفات رندة'],
        ['Teacher Name\nاسم المعلم', 'Date\nالتاريخ', 'Forecast Salary', 'Actual Salary',
         'Notes', None, 'Serial', 'Expense Name', 'Expense Date', 'Value'],
        ['Dalia', None, 4000, None, None, None, 1, 'داليا', None, -500],
        ['Total Salaries\nإجمالي الرواتب', None, 4000, 0, None, None, None,
         'Total Expenses\nإجمالي المصروفات', None, -1450],
        [],
    ]
    return rows


class SummaryLabelTests(unittest.TestCase):
    def test_two_line_labels_are_matched_on_either_part(self):
        values = excel_import._summary_values(_summary_block_rows())
        self.assertEqual(values, {
            'opening_balance': 4090.0, 'on_hand_randa': 76300.0, 'student_count': 86.0,
            'collected': 77750.0, 'expected_salaries': 39750.0, 'paid_students': 80.0,
            'randa_cash': 77750.0, 'cash_after_salaries': 36550.0, 'unpaid_students': 6.0,
            'bank_transfer': 0.0, 'salaries_paid': 0.0, 'books': 5700.0,
            'expenses': 1450.0, 'net': 76300.0, 'remaining_total': 11100.0,
        })
        # the traps: no 'salaries' from the side-table title / totals row,
        # no audit on-hand, nothing from the student rows
        self.assertNotIn('salaries', values)
        self.assertNotIn('audit_on_hand_randa', values)
        # an explicit 0 is a value, not "missing"
        self.assertIs(type(values['bank_transfer']), float)

    def test_single_part_labels_still_work(self):
        rows = [
            ['Total Received', 'الصافي', 'Students عدد الطلاب', 'Salaries', 'Delta', 'On hand Randa'],
            [100, 40, 7, 60, -5, 55],
        ]
        self.assertEqual(excel_import._summary_values(rows), {
            'collected': 100.0, 'net': 40.0, 'student_count': 7.0,
            'salaries': 60.0, 'delta': -5.0, 'on_hand_randa': 55.0,
        })
        self.assertEqual(excel_import._summary_label_parts('Net On hand Randa\nبعد الرواتب'),
                         ('net on hand randa', 'بعد الرواتب'))
        self.assertEqual(excel_import._summary_label_parts('Books Fees\nرسوم الكتب'),
                         ('books fees', 'رسوم الكتب'))
        self.assertEqual(excel_import._summary_label_parts(76300), ('', ''))

    def test_labels_below_the_summary_block_are_ignored(self):
        rows = [[None]] * 12 + [['Net\nالصافي', 'On hand Randa\nرصيد راندا'], [1, 2]]
        values = excel_import._summary_values(rows)
        self.assertNotIn('net', values)
        # legacy audit cell: an "On hand Randa" label on row index >= 10
        self.assertEqual(values, {'audit_on_hand_randa': 2.0})

    def test_metrics_prefer_explicit_cells_and_fall_back_to_rows(self):
        parsed = excel_import._parse_records
        rows = _summary_block_rows()
        summary = excel_import._summary_values(rows)
        students = [
            {'serial': 1, 'status': 'paid', 'paid': 2500.0, 'method': 'Randa Cash'},
            {'serial': 2, 'status': 'unpaid'},
            {'serial': None, 'status': 'paid', 'paid': 300.0, 'method': 'Randa Cash'},
        ]
        entries = [
            {'kind': 'salary', 'name': 'Dalia', 'amount': 0.0, 'expected': 4000.0},
            {'kind': 'expense', 'name': 'داليا', 'amount': -500.0},
            {'kind': 'expense', 'name': 'متبقي شهر سبتمبر', 'amount': 4090.0},
        ]
        explicit = excel_import._excel_month_metrics(
            {'students': students, 'entries': entries, 'summary': summary}, '2026-10')
        self.assertEqual([explicit[k] for k in CONTRACT_KEYS],
                         [4090, 77750, 77750, 0, 1450, 76300, 39750, 36550, 0, 76300,
                          86, 80, 6, 5700, 11100])
        self.assertEqual(explicit['salaries'], explicit['salaries_paid'])
        self.assertEqual(explicit['cash_in'], 4090.0)
        self.assertEqual(explicit['income'], 77750 + 4090)
        self.assertEqual(explicit['opening_source'], 'sheet')
        fallback = excel_import._excel_month_metrics(
            {'students': students, 'entries': entries, 'summary': {}}, '2026-10')
        self.assertEqual(fallback['paid_students'], 1)
        self.assertEqual(fallback['unpaid_students'], 1)
        self.assertEqual(fallback['student_count'], 2)
        self.assertEqual(fallback['collected'], 2800.0)
        self.assertEqual(fallback['expected_salaries'], 4000.0)
        self.assertEqual(fallback['salaries_paid'], 0.0)
        self.assertEqual(fallback['expenses'], -(4090 - 500))
        self.assertEqual(fallback['on_hand_randa'], 4090 + 2800 - 500)
        self.assertEqual(fallback['cash_after_salaries'], fallback['on_hand_randa'] - 4000)
        self.assertEqual(fallback['net'], fallback['collected'] - fallback['expenses'])
        self.assertEqual(fallback['opening_balance'], 0.0)     # no نسرين line, no cell
        self.assertEqual(fallback['opening_source'], '')
        self.assertIs(parsed, excel_import._parse_records)


class _FakeParams(object):
    def __init__(self):
        self.store = {}

    def sudo(self):
        return self

    def get_param(self, key, default=''):
        return self.store.get(key, default)

    def set_param(self, key, value):
        self.store[key] = value


class _FakeEnv(dict):
    def __init__(self):
        super().__init__()
        self['ir.config_parameter'] = _FakeParams()


@unittest.skipUnless(os.path.exists(REAL_WORKBOOK), 'randa.xlsx not next to the tests')
class RealWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with open(REAL_WORKBOOK, 'rb') as handle:
            raw = handle.read()
        files = [{'name': 'randa.xlsx', 'content': base64.b64encode(raw).decode('ascii')}]
        cls.parsed = {ym: excel_import._parse_records(files, ym) for ym in CONTRACT_VALUES}
        cls.metrics = {ym: excel_import._excel_month_metrics(cls.parsed[ym], ym)
                       for ym in CONTRACT_VALUES}

    def test_every_summary_cell_is_read_for_both_months(self):
        for ym in ('2026-09', '2026-10'):
            summary = self.parsed[ym]['summary']
            self.assertEqual(sorted(summary), sorted(CONTRACT_KEYS), ym)
            self.assertEqual([self.metrics[ym][key] for key in CONTRACT_KEYS],
                             list(CONTRACT_VALUES[ym]), ym)
            self.assertEqual(sorted(self.metrics[ym]['explicit_keys']), sorted(CONTRACT_KEYS))
        for ym in ('2026-09', '2026-10'):
            for key in ('student_count', 'paid_students', 'unpaid_students'):
                self.assertIs(type(self.metrics[ym][key]), int, '%s %s' % (ym, key))

    def test_legacy_chart_identity_matches_the_workbook_net(self):
        expected_income = {'2026-09': 276215 + 10360, '2026-10': 77750 + 4090}
        expected_cash_in = {'2026-09': 10360, '2026-10': 4090}
        for ym in ('2026-09', '2026-10'):
            m = self.metrics[ym]
            row = excel_import._ext_fin_row(m)
            self.assertEqual(row['income'] - row['salaries'] - row['other'], row['net'], ym)
            self.assertEqual(row['net'], CONTRACT_VALUES[ym][9], ym)
            self.assertEqual(row['income'], expected_income[ym], ym)
            self.assertEqual(row['extra_income'], expected_cash_in[ym], ym)
            self.assertEqual(m['cash_in'], expected_cash_in[ym], ym)
            self.assertEqual(row['salaries'], row['salaries_paid'], ym)
            self.assertEqual(row['student_income'], row['collected'], ym)
            for key in CONTRACT_KEYS:
                ext_key = {'randa_cash': 'cash', 'bank_transfer': 'transfer'}.get(key, key)
                self.assertEqual(row[ext_key], m[key], '%s %s' % (ym, key))
        # the first month's legacy opening (نسرين line) survives for nursery.month
        self.assertEqual(self.metrics['2026-09']['legacy_opening'], 10360.0)
        self.assertEqual(self.metrics['2026-10']['legacy_opening'], 0.0)
        self.assertEqual(self.metrics['2026-09']['opening_source'], 'sheet')

    def test_fallbacks_agree_with_the_sheet_where_the_tables_allow(self):
        # student_count is deliberately not asserted here: the sheet's G3 sums the
        # class counters (Oct: 86) while the numbered-row fallback gives 85 because
        # two numbered rows share a name and merge; the explicit cell wins anyway.
        for ym in ('2026-09', '2026-10'):
            parsed = dict(self.parsed[ym], summary={})
            m = excel_import._excel_month_metrics(parsed, ym)
            self.assertEqual(m['salaries_paid'], CONTRACT_VALUES[ym][8], ym)
            self.assertEqual(m['expected_salaries'], 39750.0, ym)
            self.assertEqual(m['on_hand_randa'], CONTRACT_VALUES[ym][5], ym)
            self.assertEqual(m['net'], CONTRACT_VALUES[ym][9], ym)
            self.assertEqual(m['expenses'], CONTRACT_VALUES[ym][4], ym)
            self.assertEqual(m['books'], CONTRACT_VALUES[ym][13], ym)
            self.assertEqual(m['remaining_total'], CONTRACT_VALUES[ym][14], ym)
            self.assertEqual(m['unpaid_students'], CONTRACT_VALUES[ym][12], ym)
        # The paid/unpaid fallback counts numbered rows by their Payment status
        # (the merged duplicate above makes it one short of the sheet's G5).
        for ym in ('2026-09', '2026-10'):
            m = excel_import._excel_month_metrics(dict(self.parsed[ym], summary={}), ym)
            numbered = [s for s in self.parsed[ym]['students'] if s.get('serial') is not None]
            self.assertEqual(m['paid_students'],
                             sum(1 for s in numbered if s.get('status') == 'paid'), ym)
            self.assertEqual(m['unpaid_students'],
                             sum(1 for s in numbered if s.get('status') == 'unpaid'), ym)
            self.assertEqual(m['student_count'], len(numbered), ym)

    def test_empty_month_parses(self):
        m = self.metrics['2026-11']
        self.assertEqual([m[key] for key in CONTRACT_KEYS], list(CONTRACT_VALUES['2026-11']))
        self.assertEqual(self.parsed['2026-11']['warnings'], [w for w in self.parsed['2026-11']['warnings']
                                                            if w.startswith('تم تجاهل شيت')])

    def test_refresh_writes_only_the_two_params(self):
        env = _FakeEnv()
        env['ir.config_parameter'].store['nursery.ext_fin'] = json.dumps(
            {'2026-09': {'ym': '2026-09', 'income': 1.0}})
        metrics = excel_import.refresh_month_summary(self.parsed['2026-10'], '2026-10', env)
        store = env['ir.config_parameter'].store
        self.assertEqual(sorted(store), ['nursery.excel_month_2026-10', 'nursery.ext_fin'])
        self.assertEqual(json.loads(store['nursery.excel_month_2026-10']),
                         excel_import._excel_month_payload(metrics))
        # October states its own positive opening (A3 = 4,090): nursery.month reads it as-is
        self.assertEqual(json.loads(store['nursery.excel_month_2026-10'])['opening_balance'], 4090.0)
        ext = json.loads(store['nursery.ext_fin'])
        self.assertEqual(ext['2026-09'], {'ym': '2026-09', 'income': 1.0})
        self.assertEqual(ext['2026-10'], excel_import._ext_fin_row(metrics))
        self.assertEqual([ext['2026-10'][k] for k in (
            'opening_balance', 'collected', 'cash', 'transfer', 'expenses', 'on_hand_randa',
            'expected_salaries', 'cash_after_salaries', 'salaries_paid', 'net', 'student_count',
            'paid_students', 'unpaid_students', 'books', 'remaining_total')],
            list(CONTRACT_VALUES['2026-10']))

    def test_month_opening_keeps_the_legacy_notion_for_nursery_month(self):
        # الكروت: خلية A3 الصريحة (سبتمبر 0 معتمد). nursery.month: المعنى القديم
        # (سطر نسرين الموجب لأول شهر، وإلا A3 الموجب)، والصفر عنده = «يرث».
        sep, oct_, nov = (self.metrics[ym] for ym in ('2026-09', '2026-10', '2026-11'))
        self.assertEqual(sep['opening_balance'], 0.0)
        self.assertEqual(sep['opening_source'], 'sheet')
        self.assertEqual(sep['month_opening'], 10360.0)
        self.assertEqual(sep['month_opening_source'], 'نسرين')
        self.assertEqual(oct_['opening_balance'], 4090.0)
        self.assertEqual(oct_['month_opening'], 4090.0)
        self.assertEqual(oct_['month_opening_source'], 'sheet')
        self.assertEqual(nov['month_opening'], 0.0)
        self.assertEqual(nov['month_opening_source'], '')
        for ym, metrics in (('2026-09', sep), ('2026-10', oct_), ('2026-11', nov)):
            payload = excel_import._excel_month_payload(metrics)
            row = excel_import._ext_fin_row(metrics)
            # ext_fin (cards) carries the explicit cell; the month payload the site notion
            self.assertEqual(row['opening_balance'], metrics['opening_balance'], ym)
            self.assertEqual(payload['opening_balance'], metrics['month_opening'], ym)
            self.assertEqual(payload['opening_source'], metrics['month_opening_source'], ym)
            self.assertEqual(payload['sheet_opening_balance'], metrics['opening_balance'], ym)
            self.assertEqual(payload['sheet_opening_source'], metrics['opening_source'], ym)
            # everything else is untouched
            for key, value in metrics.items():
                if key not in ('opening_balance', 'opening_source'):
                    self.assertEqual(payload[key], value, '%s %s' % (ym, key))

    def test_refresh_september_keeps_nesrin_as_the_site_opening(self):
        env = _FakeEnv()
        excel_import.refresh_month_summary(self.parsed['2026-09'], '2026-09', env)
        store = env['ir.config_parameter'].store
        excel = json.loads(store['nursery.excel_month_2026-09'])
        # exactly what models_nursery._sheet_opening() computes from the param
        sheet_opening = float(excel.get('opening_balance') or 0.0) if excel else 0.0
        self.assertEqual(sheet_opening, 10360.0)       # > 0 → September never inherits August
        self.assertEqual(excel['sheet_opening_balance'], 0.0)
        self.assertEqual(excel['opening_source'], 'نسرين')
        # models_nursery._carry_closing(): opening + randa_cash + expenses_paid_out == D3
        self.assertEqual(sheet_opening + excel['randa_cash'] + excel['expenses_paid_out'], 4090.0)
        # the dashboard row still shows the sheet's explicit 0
        ext = json.loads(store['nursery.ext_fin'])['2026-09']
        self.assertEqual(ext['opening_balance'], 0.0)
        self.assertEqual(ext['opening_source'], 'sheet')

    def test_month_record_opening_rule(self):
        august = types.SimpleNamespace(opening_balance=1.0, _carry_closing=lambda: 99999.0)
        # September: A3 = 0 but the Nesrin line is the site's opening — not August's closing
        self.assertEqual(excel_import._month_record_opening(self.metrics['2026-09'], august), 10360.0)
        self.assertEqual(excel_import._month_record_opening(self.metrics['2026-09'], None), 10360.0)
        # October: explicit A3 > 0 pins it
        self.assertEqual(excel_import._month_record_opening(self.metrics['2026-10'], august), 4090.0)
        # no opening of its own → previous month's carried closing, else 0
        self.assertEqual(excel_import._month_record_opening(self.metrics['2026-11'], august), 99999.0)
        self.assertEqual(excel_import._month_record_opening(self.metrics['2026-11'], None), 0.0)
        legacy_prev = types.SimpleNamespace(opening_balance=7.0, _month_closing=lambda: 42.0)
        self.assertEqual(excel_import._month_record_opening({'month_opening': 0.0}, legacy_prev), 42.0)
        bare_prev = types.SimpleNamespace(opening_balance=7.0)
        self.assertEqual(excel_import._month_record_opening({'month_opening': 0.0}, bare_prev), 7.0)


if __name__ == '__main__':
    unittest.main()
